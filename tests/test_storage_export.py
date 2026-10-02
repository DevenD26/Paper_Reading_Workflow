import json
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import fitz
from PIL import Image
from docx import Document

from paper_notes.exporters import (
    build_markdown,
    export_combined_papers,
    export_combined_summaries,
    export_docx,
    export_markdown,
    safe_export_stem,
)
from paper_notes.pdf_tools import crop_page, page_count, render_page
from paper_notes.storage import (
    PAPER_ORDER_FILENAME,
    blank_data_source,
    blank_empirical_strategy,
    blank_exhibit,
    create_record,
    delete_record,
    list_records,
    load_notes,
    move_record,
    normalize_notes,
    safe_slug,
    save_notes,
    save_uploaded_image,
    section_completion,
    update_record_title,
    validate_pdf_upload,
)
from paper_notes.tracker import TRACKER_COLUMNS, build_tracker_row, build_tracker_rows


_minimal_document = fitz.open()
_minimal_document.new_page()
MINIMAL_PDF = _minimal_document.tobytes()
_minimal_document.close()


class StorageAndExportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name) / "papers"

    def tearDown(self):
        self.temporary.cleanup()

    def test_safe_slug_and_unique_record_folders(self):
        self.assertEqual(safe_slug("  A Paper: Caf\u00e9 & Politics!  "), "a-paper-cafe-politics")
        first = create_record(self.root, "Same Paper", MINIMAL_PDF, "original.pdf")
        second = create_record(self.root, "Same Paper", MINIMAL_PDF, "original.pdf")
        self.assertEqual(first["record_id"], "same-paper")
        self.assertEqual(second["record_id"], "same-paper-2")
        self.assertTrue((self.root / "same-paper" / "source.pdf").is_file())
        self.assertEqual(
            [item["record_id"] for item in list_records(self.root)],
            [second["record_id"], first["record_id"]],
        )

    def test_manual_moves_persist_and_stop_at_boundaries(self):
        first = create_record(self.root, "First", MINIMAL_PDF, "first.pdf")
        second = create_record(self.root, "Second", MINIMAL_PDF, "second.pdf")
        third = create_record(self.root, "Third", MINIMAL_PDF, "third.pdf")

        self.assertEqual(
            [item["record_id"] for item in list_records(self.root)],
            [third["record_id"], second["record_id"], first["record_id"]],
        )
        self.assertTrue(move_record(self.root, second["record_id"], -1))
        expected = [second["record_id"], third["record_id"], first["record_id"]]
        self.assertEqual([item["record_id"] for item in list_records(self.root)], expected)
        self.assertEqual([item["record_id"] for item in list_records(self.root)], expected)
        self.assertFalse(move_record(self.root, second["record_id"], -1))
        self.assertFalse(move_record(self.root, first["record_id"], 1))
        self.assertEqual([item["record_id"] for item in list_records(self.root)], expected)

    def test_legacy_order_migration_is_deterministic_and_does_not_rewrite_notes(self):
        oldest = create_record(self.root, "Oldest", MINIMAL_PDF, "oldest.pdf")
        newest = create_record(self.root, "Newest", MINIMAL_PDF, "newest.pdf")
        timestamps = {
            oldest["record_id"]: "2026-01-01T00:00:00+00:00",
            newest["record_id"]: "2026-02-01T00:00:00+00:00",
        }
        notes_paths = []
        for record_id, timestamp in timestamps.items():
            path = self.root / record_id / "notes.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["updated_at"] = timestamp
            path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
            notes_paths.append(path)
        (self.root / PAPER_ORDER_FILENAME).unlink()
        before = {path: path.read_bytes() for path in notes_paths}

        self.assertEqual(
            [item["record_id"] for item in list_records(self.root)],
            [newest["record_id"], oldest["record_id"]],
        )
        self.assertEqual(before, {path: path.read_bytes() for path in notes_paths})
        self.assertTrue((self.root / PAPER_ORDER_FILENAME).is_file())

    def test_new_papers_start_on_top_and_edits_do_not_reorder(self):
        first = create_record(self.root, "First", MINIMAL_PDF, "first.pdf")
        second = create_record(self.root, "Second", MINIMAL_PDF, "second.pdf")
        move_record(self.root, first["record_id"], -1)
        edited = load_notes(self.root, second["record_id"])
        edited["title"] = "Renamed Second"
        edited["categories"] = ["Week 6", "Methods"]
        edited["created_at"] = "1999-01-01T00:00:00+00:00"
        edited["sections"]["additional_notes"] = "Exact authored text."
        save_notes(self.root, edited)
        self.assertEqual(
            [item["record_id"] for item in list_records(self.root)],
            [first["record_id"], second["record_id"]],
        )

        third = create_record(self.root, "Third", MINIMAL_PDF, "third.pdf")
        self.assertEqual(
            [item["record_id"] for item in list_records(self.root)],
            [third["record_id"], first["record_id"], second["record_id"]],
        )

    def test_invalid_and_stale_order_metadata_is_repaired(self):
        first = create_record(self.root, "First", MINIMAL_PDF, "first.pdf")
        second = create_record(self.root, "Second", MINIMAL_PDF, "second.pdf")
        third = create_record(self.root, "Third", MINIMAL_PDF, "third.pdf")
        order_path = self.root / PAPER_ORDER_FILENAME
        order_path.write_text(
            json.dumps(
                {
                    "schema_version": 1,
                    "record_ids": [
                        first["record_id"],
                        "deleted-record",
                        first["record_id"],
                        42,
                        third["record_id"],
                    ],
                }
            ),
            encoding="utf-8",
        )

        expected = [second["record_id"], first["record_id"], third["record_id"]]
        self.assertEqual([item["record_id"] for item in list_records(self.root)], expected)
        repaired = json.loads(order_path.read_text(encoding="utf-8"))
        self.assertEqual(repaired["record_ids"], expected)

        order_path.write_text("{not valid json", encoding="utf-8")
        fallback = [third["record_id"], second["record_id"], first["record_id"]]
        self.assertEqual([item["record_id"] for item in list_records(self.root)], fallback)
        self.assertEqual(
            json.loads(order_path.read_text(encoding="utf-8"))["record_ids"], fallback
        )

    def test_notes_round_trip_preserves_wording_and_unicode(self):
        notes = create_record(self.root, "My Paper", MINIMAL_PDF, "paper.pdf")
        authored = "- First bullet\n- Exact phrasing: \u03b2 = 0.42\n\nA final paragraph."
        notes["sections"]["one_minute_summary"] = authored
        saved = save_notes(self.root, notes)
        loaded = load_notes(self.root, notes["record_id"])
        self.assertEqual(loaded["sections"]["one_minute_summary"], authored)
        self.assertEqual(saved["updated_at"], loaded["updated_at"])
        self.assertEqual(len(list_records(self.root)), 1)

    def test_title_edit_preserves_record_folder_and_other_notes(self):
        notes = create_record(self.root, "Original Title", MINIMAL_PDF, "paper.pdf")
        notes["sections"]["one_minute_summary"] = "Keep this exact note."
        save_notes(self.root, notes)
        folder = self.root / notes["record_id"]
        updated = update_record_title(self.root, notes["record_id"], "Edited Title")
        self.assertEqual(updated["title"], "Edited Title")
        self.assertEqual(updated["sections"]["one_minute_summary"], "Keep this exact note.")
        self.assertTrue(folder.is_dir())
        self.assertEqual(updated["record_id"], notes["record_id"])
        with self.assertRaises(ValueError):
            update_record_title(self.root, notes["record_id"], "   ")

    def test_delete_record_removes_entire_paper_folder_only(self):
        deleted = create_record(self.root, "Delete Me", MINIMAL_PDF, "delete.pdf")
        retained = create_record(self.root, "Keep Me", MINIMAL_PDF, "keep.pdf")
        deleted_folder = self.root / deleted["record_id"]
        formatted_note = deleted_folder / "formatted_notes" / "Delete Me_notes.docx"
        formatted_note.write_bytes(b"saved export")
        screenshot = deleted_folder / "exhibits" / "saved-screenshot.png"
        screenshot.write_bytes(b"saved image")

        removed_notes = delete_record(self.root, deleted["record_id"])

        self.assertEqual(removed_notes["title"], "Delete Me")
        self.assertFalse(deleted_folder.exists())
        self.assertTrue((self.root / retained["record_id"]).is_dir())
        self.assertEqual([item["record_id"] for item in list_records(self.root)], [retained["record_id"]])
        self.assertEqual(
            json.loads((self.root / PAPER_ORDER_FILENAME).read_text(encoding="utf-8"))["record_ids"],
            [retained["record_id"]],
        )

    def test_delete_record_rejects_invalid_identifier(self):
        with self.assertRaises(ValueError):
            delete_record(self.root, "../outside")

    def test_completion_only_reports_user_content(self):
        notes = create_record(self.root, "Status Test", MINIMAL_PDF, "paper.pdf")
        self.assertFalse(any(section_completion(notes).values()))
        source = blank_data_source()
        source["description"] = "Census microdata"
        notes["data_sources"].append(source)
        exhibit = blank_exhibit()
        exhibit["identifier_title"] = "Figure 1"
        notes["exhibits"].append(exhibit)
        status = section_completion(notes)
        self.assertTrue(status["data_sources"])
        self.assertTrue(status["exhibits"])
        self.assertFalse(status["contribution"])

    def test_markdown_and_docx_exports_are_created(self):
        notes = create_record(self.root, "Export Test", MINIMAL_PDF, "paper.pdf")
        exact_text = "- User bullet\nSecond line stays here."
        notes["sections"]["one_minute_summary"] = exact_text
        source = blank_data_source()
        source.update({"name": "Administrative records", "description": "Outcomes for 2010–2020."})
        notes["data_sources"].append(source)
        exhibit = blank_exhibit()
        exhibit.update({"identifier_title": "Table 3", "page_number": 14, "interpretation": "My reading."})
        notes["exhibits"].append(exhibit)
        markdown = build_markdown(notes)
        self.assertIn(exact_text, markdown)
        self.assertIn("## Tables and Figures", markdown)
        self.assertIn("### Table 3", markdown)
        self.assertNotIn("Not yet completed", markdown)
        self.assertNotIn("## Contribution to the Literature", markdown)
        markdown_path = export_markdown(self.root, notes)
        docx_path = export_docx(self.root, notes)
        self.assertTrue(markdown_path.is_file())
        self.assertTrue(docx_path.is_file())
        self.assertEqual(markdown_path.parent.name, "formatted_notes")
        self.assertEqual(markdown_path.name, "Export Test_notes.md")
        self.assertEqual(docx_path.name, "Export Test_notes.docx")
        self.assertEqual(markdown_path.read_text(encoding="utf-8"), markdown)

    def test_legacy_notes_gain_optional_fields_without_losing_authored_text(self):
        legacy = {
            "schema_version": 1,
            "record_id": "legacy",
            "title": "Legacy",
            "sections": {"one_minute_summary": "Exact existing note"},
            "data_sources": [{"id": "source", "name": "Existing source", "description": "Kept"}],
            "exhibits": [],
        }
        migrated = normalize_notes(legacy)
        self.assertEqual(migrated["sections"]["one_minute_summary"], "Exact existing note")
        self.assertEqual(migrated["data_sources"], legacy["data_sources"])
        self.assertEqual(migrated["sections"]["additional_notes"], "")
        self.assertEqual(migrated["empirical_strategies"], [])
        self.assertEqual(migrated["categories"], [])
        self.assertNotIn("additional_notes", legacy["sections"])

    def test_empty_sections_and_fields_are_omitted_from_markdown(self):
        notes = create_record(self.root, "Mostly Blank", MINIMAL_PDF, "paper.pdf")
        notes["sections"]["additional_notes"] = "A miscellaneous observation."
        markdown = build_markdown(notes)
        self.assertIn("## Additional Notes", markdown)
        self.assertNotIn("## Data Sources", markdown)
        self.assertNotIn("## Tables and Figures", markdown)
        self.assertNotIn("Not yet completed", markdown)

    def test_word_export_places_exhibit_image_before_notes(self):
        notes = create_record(self.root, "Image Order", MINIMAL_PDF, "paper.pdf")
        image_path = self.root / notes["record_id"] / "exhibits" / "ordering.png"
        Image.new("RGB", (40, 30), "white").save(image_path)
        exhibit = blank_exhibit()
        exhibit.update(
            {
                "identifier_title": "Table image first",
                "interpretation": "Interpretation comes after the picture.",
                "image": "exhibits/ordering.png",
            }
        )
        notes["exhibits"].append(exhibit)
        output = export_docx(self.root, notes)
        with zipfile.ZipFile(output) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        self.assertLess(xml.index("Table image first"), xml.index("w:drawing"))
        self.assertLess(xml.index("w:drawing"), xml.index("Interpretation comes after the picture."))

    def test_safe_export_filename_keeps_title_and_removes_path_characters(self):
        self.assertEqual(safe_export_stem("Paper / Lobbying: Evidence"), "Paper _ Lobbying_ Evidence_notes")

    def test_combined_summary_export_includes_selected_authored_summaries(self):
        first = create_record(self.root, "First Paper", MINIMAL_PDF, "first.pdf")
        first["sections"]["one_minute_summary"] = "My first authored summary."
        first["categories"] = ["Week 6"]
        second = create_record(self.root, "Second Paper", MINIMAL_PDF, "second.pdf")
        second["sections"]["one_minute_summary"] = "My second authored summary."
        blank = create_record(self.root, "Blank Paper", MINIMAL_PDF, "blank.pdf")

        output = export_combined_summaries(
            self.root,
            [first, second, blank],
            "Week 6 summaries",
        )
        self.assertEqual(output.parent.name, "combined_exports")
        self.assertEqual(output.name, "Week 6 summaries.docx")
        text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
        self.assertIn("First Paper", text)
        self.assertIn("My first authored summary.", text)
        self.assertIn("Categories: Week 6", text)
        self.assertIn("Second Paper", text)
        self.assertNotIn("Blank Paper", text)

        second_output = export_combined_summaries(self.root, [first], "Week 6 summaries")
        self.assertEqual(second_output.name, "Week 6 summaries_2.docx")

    def test_combined_export_follows_custom_order_not_argument_or_title_order(self):
        alpha = create_record(self.root, "Duplicate", MINIMAL_PDF, "alpha.pdf")
        beta = create_record(self.root, "Duplicate", MINIMAL_PDF, "beta.pdf")
        gamma = create_record(self.root, "Gamma", MINIMAL_PDF, "gamma.pdf")
        for notes, marker in [(alpha, "ALPHA MARKER"), (beta, "BETA MARKER"), (gamma, "GAMMA MARKER")]:
            notes["sections"]["one_minute_summary"] = marker
            save_notes(self.root, notes)
        move_record(self.root, alpha["record_id"], -1)
        expected_ids = [gamma["record_id"], alpha["record_id"], beta["record_id"]]
        self.assertEqual([item["record_id"] for item in list_records(self.root)], expected_ids)

        output = export_combined_summaries(
            self.root,
            [beta, alpha, gamma],
            "Ordered summaries",
        )
        text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
        self.assertLess(text.index("GAMMA MARKER"), text.index("ALPHA MARKER"))
        self.assertLess(text.index("ALPHA MARKER"), text.index("BETA MARKER"))

    def test_group_export_selected_sections_order_columns_and_arial(self):
        notes = create_record(self.root, "Full Paper", MINIMAL_PDF, "paper.pdf")
        notes["sections"].update(
            {
                "one_minute_summary": "SUMMARY MARKER",
                "contribution": "CONTRIBUTION MARKER",
                "background": "BACKGROUND SHOULD BE OMITTED",
                "empirical_strategy": "EMPIRICAL MARKER",
            }
        )
        strategy = blank_empirical_strategy()
        strategy["description"] = "STRATEGY ENTRY MARKER"
        notes["empirical_strategies"].append(strategy)
        exhibit = blank_exhibit()
        exhibit["identifier_title"] = "FIGURE MARKER"
        exhibit["interpretation"] = "FIGURE NOTES MARKER"
        notes["exhibits"].append(exhibit)

        output = export_combined_papers(
            self.root,
            [notes],
            "Selected Review",
            ["one_minute_summary", "contribution", "empirical_strategy", "exhibits"],
        )
        document = Document(output)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("SUMMARY MARKER", text)
        self.assertIn("CONTRIBUTION MARKER", text)
        self.assertIn("EMPIRICAL MARKER", text)
        self.assertIn("FIGURE MARKER", text)
        self.assertNotIn("BACKGROUND SHOULD BE OMITTED", text)
        self.assertLess(text.index("SUMMARY MARKER"), text.index("EMPIRICAL MARKER"))
        self.assertLess(text.index("EMPIRICAL MARKER"), text.index("FIGURE MARKER"))

        with zipfile.ZipFile(output) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
            styles_xml = archive.read("word/styles.xml").decode("utf-8")
        self.assertIn('w:num="2"', document_xml)
        self.assertIn('w:num="1"', document_xml)
        self.assertIn('w:ascii="Arial"', styles_xml)
        title_style = styles_xml.split('w:styleId="Title"', 1)[1].split("</w:style>", 1)[0]
        self.assertNotIn("w:pBdr", title_style)
        self.assertIn('w:color w:val="000000"', title_style)

    def test_individual_word_export_uses_arial(self):
        notes = create_record(self.root, "Arial Paper", MINIMAL_PDF, "paper.pdf")
        notes["sections"]["one_minute_summary"] = "Arial body text"
        output = export_docx(self.root, notes)
        with zipfile.ZipFile(output) as archive:
            styles_xml = archive.read("word/styles.xml").decode("utf-8")
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('w:ascii="Arial"', styles_xml)
        self.assertIn('w:ascii="Arial"', document_xml)

    def test_individual_word_export_uses_review_first_two_column_layout(self):
        notes = create_record(self.root, "Structured Paper", MINIMAL_PDF, "paper.pdf")
        notes["sections"].update(
            {
                "one_minute_summary": "SUMMARY MARKER",
                "concerns_extensions": "CONCERNS MARKER",
                "empirical_strategy": "EMPIRICAL MARKER",
            }
        )
        exhibit = blank_exhibit()
        exhibit.update(
            {
                "identifier_title": "FIGURE MARKER",
                "interpretation": "FIGURE NOTES MARKER",
            }
        )
        notes["exhibits"].append(exhibit)

        output = export_docx(self.root, notes)
        document = Document(output)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertLess(text.index("SUMMARY MARKER"), text.index("CONCERNS MARKER"))
        self.assertLess(text.index("CONCERNS MARKER"), text.index("EMPIRICAL MARKER"))
        self.assertLess(text.index("EMPIRICAL MARKER"), text.index("FIGURE MARKER"))

        with zipfile.ZipFile(output) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('w:num="2"', document_xml)
        self.assertIn('w:num="1"', document_xml)

    def test_tracker_mapping_is_stable_and_spreadsheet_ready(self):
        notes = create_record(self.root, "Tracker Paper", MINIMAL_PDF, "paper.pdf")
        notes["categories"] = ["Week 6", "Lobbying"]
        notes["sections"]["one_minute_summary"] = "Tracker summary"
        notes["sections"]["background"] = "Tracker background"
        notes["sections"]["concerns_extensions"] = "Tracker concerns"
        source = blank_data_source()
        source.update({"name": "Administrative records", "description": "Outcomes"})
        notes["data_sources"].append(source)
        row = build_tracker_row(notes)
        self.assertEqual(list(row), TRACKER_COLUMNS)
        self.assertEqual(row["record_id"], notes["record_id"])
        self.assertEqual(row["categories"], "Week 6, Lobbying")
        self.assertEqual(row["data_sources"], "Administrative records: Outcomes")
        updated = dict(notes)
        updated["title"] = "Updated Tracker Paper"
        rows = build_tracker_rows([notes, updated])
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["title"], "Updated Tracker Paper")

    def test_invalid_upload_is_rejected_without_folder(self):
        with self.assertRaises(ValueError):
            create_record(self.root, "Not PDF", b"plain text", "notes.txt")
        self.assertFalse(self.root.exists())

    def test_unreadable_and_oversized_uploads_leave_no_partial_record(self):
        with self.assertRaises(ValueError):
            create_record(self.root, "Broken", b"%PDF-1.4\nnot a pdf", "broken.pdf")
        self.assertFalse(self.root.exists())
        with patch("paper_notes.storage.MAX_PDF_UPLOAD_BYTES", 10):
            with self.assertRaises(ValueError):
                validate_pdf_upload(MINIMAL_PDF)
        self.assertFalse(self.root.exists())

    def test_record_layout_contains_one_pdf_and_all_unified_folders(self):
        notes = create_record(self.root, "Hub paper", MINIMAL_PDF, "paper.pdf")
        folder = self.root / notes["record_id"]
        self.assertEqual(list(folder.rglob("*.pdf")), [folder / "source.pdf"])
        self.assertEqual((folder / "source.pdf").read_bytes(), MINIMAL_PDF)
        for name in ("exhibits", "extraction_proposals", "formatted_notes"):
            self.assertTrue((folder / name).is_dir())
        self.assertEqual(load_notes(self.root, notes["record_id"])["sections"]["one_minute_summary"], "")

    def test_pdf_page_render_and_manual_crop(self):
        pdf_path = Path(self.temporary.name) / "crop-source.pdf"
        document = fitz.open()
        page = document.new_page(width=400, height=300)
        page.insert_text((40, 50), "Exhibit area")
        document.save(pdf_path)
        document.close()

        self.assertEqual(page_count(pdf_path), 1)
        rendered = render_page(pdf_path, 1, zoom=1)
        self.assertEqual(rendered.size, (400, 300))
        output = Path(self.temporary.name) / "exhibit.png"
        crop_page(pdf_path, 1, (10, 20, 90, 80), output, zoom=1)
        self.assertTrue(output.is_file())
        self.assertEqual(render_page(pdf_path, 1, zoom=1).mode, "RGB")

    def test_uploaded_screenshot_is_saved_as_record_image(self):
        notes = create_record(self.root, "Screenshot Upload", MINIMAL_PDF, "paper.pdf")
        exhibit = blank_exhibit()
        screenshot = BytesIO()
        Image.new("RGB", (120, 80), "blue").save(screenshot, format="JPEG")

        relative_path = save_uploaded_image(
            self.root,
            notes["record_id"],
            "exhibit",
            exhibit["id"],
            screenshot.getvalue(),
        )

        self.assertEqual(relative_path, f"exhibits/exhibit-{exhibit['id']}.png")
        saved_path = self.root / notes["record_id"] / relative_path
        self.assertTrue(saved_path.is_file())
        with Image.open(saved_path) as saved:
            self.assertEqual(saved.format, "PNG")
            self.assertEqual(saved.size, (120, 80))

    def test_invalid_screenshot_upload_is_rejected(self):
        notes = create_record(self.root, "Bad Screenshot", MINIMAL_PDF, "paper.pdf")
        strategy = blank_empirical_strategy()
        with self.assertRaisesRegex(ValueError, "could not be read as an image"):
            save_uploaded_image(
                self.root,
                notes["record_id"],
                "strategy",
                strategy["id"],
                b"not an image",
            )
        expected = self.root / notes["record_id"] / "exhibits" / f"strategy-{strategy['id']}.png"
        self.assertFalse(expected.exists())


if __name__ == "__main__":
    unittest.main()

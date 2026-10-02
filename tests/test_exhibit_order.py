import hashlib
import os
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from PIL import Image
from streamlit.testing.v1 import AppTest

from paper_notes.exporters import build_markdown, export_combined_papers, export_docx
from paper_notes.storage import (
    blank_exhibit,
    create_record,
    load_notes,
    move_exhibit,
    mutate_notes,
    save_notes,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
_minimal_document = fitz.open()
_minimal_document.new_page()
MINIMAL_PDF = _minimal_document.tobytes()
_minimal_document.close()


def _titles(notes):
    return [exhibit["identifier_title"] for exhibit in notes["exhibits"]]


def _sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class ExhibitOrderTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name) / "papers"
        self.notes = create_record(self.root, "Ordered exhibits", MINIMAL_PDF, "paper.pdf")
        self.folder = self.root / self.notes["record_id"]

    def tearDown(self):
        self.temporary.cleanup()

    def _exhibit(self, title, color, *, extracted=False):
        exhibit = blank_exhibit()
        image_name = f"{exhibit['id']}.png"
        image_path = self.folder / "exhibits" / image_name
        Image.new("RGB", (12, 8), color).save(image_path)
        exhibit.update(
            {
                "identifier_title": title,
                "page_number": 7,
                "convey": f"{title} convey exact text",
                "empirical_strategy": f"{title} strategy exact text",
                "interpretation": f"{title} interpretation exact text",
                "qualifications": f"{title} qualifications exact text",
                "image": f"exhibits/{image_name}",
                "crop": {"page_number": 7, "left": 1.5, "top": 2.5, "right": 98.5, "bottom": 97.5},
                "images": [
                    {
                        "label": "panel_A",
                        "page_number": 7,
                        "image": f"exhibits/{image_name}",
                        "crop": {
                            "page_number": 7,
                            "left": 3.0,
                            "top": 4.0,
                            "right": 96.0,
                            "bottom": 95.0,
                        },
                    }
                ],
            }
        )
        if extracted:
            exhibit["extraction_source"] = {
                "source_sha256": "source-checksum",
                "candidate_id": f"candidate-{exhibit['id']}",
                "caption_text": f"Caption for {title}",
                "proposal_id": "proposal-1",
            }
        return exhibit

    def test_moves_persist_and_preserve_every_exhibit_field_and_image_checksum(self):
        manual = self._exhibit("Manual screenshot", "red")
        extracted = self._exhibit("Extractor result", "blue", extracted=True)
        last = self._exhibit("Later manual", "green")
        self.notes["exhibits"] = [manual, extracted, last]
        saved = save_notes(self.root, self.notes)
        originals = {item["id"]: deepcopy(item) for item in saved["exhibits"]}
        checksums = {
            item["id"]: _sha256(self.folder / item["image"])
            for item in saved["exhibits"]
        }

        moved_down = move_exhibit(self.root, self.notes["record_id"], manual["id"], 1)
        self.assertEqual(_titles(moved_down), ["Extractor result", "Manual screenshot", "Later manual"])
        moved_up = move_exhibit(self.root, self.notes["record_id"], last["id"], -1)
        self.assertEqual(_titles(moved_up), ["Extractor result", "Later manual", "Manual screenshot"])
        reloaded = load_notes(self.root, self.notes["record_id"])
        self.assertEqual(_titles(reloaded), ["Extractor result", "Later manual", "Manual screenshot"])
        self.assertEqual({item["id"]: item for item in reloaded["exhibits"]}, originals)
        self.assertEqual(
            {item["id"]: _sha256(self.folder / item["image"]) for item in reloaded["exhibits"]},
            checksums,
        )

    def test_first_and_last_boundaries_are_safe_no_ops(self):
        first = self._exhibit("First", "red")
        last = self._exhibit("Last", "blue")
        self.notes["exhibits"] = [first, last]
        saved = save_notes(self.root, self.notes)

        unchanged_first = move_exhibit(self.root, self.notes["record_id"], first["id"], -1)
        self.assertEqual(unchanged_first, saved)
        unchanged_last = move_exhibit(self.root, self.notes["record_id"], last["id"], 1)
        self.assertEqual(unchanged_last, saved)
        with self.assertRaisesRegex(ValueError, "direction"):
            move_exhibit(self.root, self.notes["record_id"], first["id"], 2)

    def test_stale_save_keeps_latest_order_and_appends_new_exhibits(self):
        first = self._exhibit("First", "red")
        second = self._exhibit("Second", "blue", extracted=True)
        self.notes["exhibits"] = [first, second]
        save_notes(self.root, self.notes)
        stale_organizer = load_notes(self.root, self.notes["record_id"])

        move_exhibit(self.root, self.notes["record_id"], second["id"], -1)
        extractor_addition = self._exhibit("New extractor result", "green", extracted=True)
        mutate_notes(
            self.root,
            self.notes["record_id"],
            lambda latest: latest["exhibits"].append(extractor_addition),
        )
        stale_organizer["exhibits"][0]["interpretation"] = "Edited without changing position"
        manual_addition = self._exhibit("New manual screenshot", "yellow")
        stale_organizer["exhibits"].append(manual_addition)
        save_notes(self.root, stale_organizer)

        reloaded = load_notes(self.root, self.notes["record_id"])
        self.assertEqual(
            _titles(reloaded),
            ["Second", "First", "New extractor result", "New manual screenshot"],
        )
        first_reloaded = next(item for item in reloaded["exhibits"] if item["id"] == first["id"])
        self.assertEqual(first_reloaded["interpretation"], "Edited without changing position")

    def test_saved_order_controls_markdown_individual_word_and_combined_word_exports(self):
        first = self._exhibit("Export first", "red")
        second = self._exhibit("Export second", "blue", extracted=True)
        self.notes["exhibits"] = [first, second]
        save_notes(self.root, self.notes)
        move_exhibit(self.root, self.notes["record_id"], second["id"], -1)
        ordered = load_notes(self.root, self.notes["record_id"])

        markdown = build_markdown(ordered)
        self.assertLess(markdown.index("Export second"), markdown.index("Export first"))
        individual = Document(export_docx(self.root, ordered))
        individual_text = "\n".join(paragraph.text for paragraph in individual.paragraphs)
        self.assertLess(individual_text.index("Export second"), individual_text.index("Export first"))
        combined = Document(
            export_combined_papers(self.root, [ordered], "Ordered combined", ["exhibits"])
        )
        combined_text = "\n".join(paragraph.text for paragraph in combined.paragraphs)
        self.assertLess(combined_text.index("Export second"), combined_text.index("Export first"))

    def test_app_move_controls_apply_immediately_and_disable_boundaries(self):
        first = self._exhibit("UI first", "red")
        second = self._exhibit("UI second", "blue", extracted=True)
        self.notes["exhibits"] = [first, second]
        save_notes(self.root, self.notes)

        with patch.dict(os.environ, {"PAPER_NOTES_DATA_ROOT": str(self.root)}):
            app = AppTest.from_file(PROJECT_ROOT / "app.py").run(timeout=30)
            self.assertEqual(len(app.exception), 0)
            up_buttons = [button for button in app.button if button.label == "↑"]
            down_buttons = [button for button in app.button if button.label == "↓"]
            self.assertEqual([button.disabled for button in up_buttons], [True, False])
            self.assertEqual([button.disabled for button in down_buttons], [False, True])
            down_buttons[0].click().run(timeout=30)

        self.assertEqual(len(app.exception), 0)
        self.assertEqual(_titles(load_notes(self.root, self.notes["record_id"])), ["UI second", "UI first"])


if __name__ == "__main__":
    unittest.main()

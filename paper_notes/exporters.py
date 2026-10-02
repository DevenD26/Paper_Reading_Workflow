from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .storage import SECTION_LABELS, order_records, record_dir


GROUP_EXPORT_SECTIONS = [
    ("one_minute_summary", "1 Minute Summary"),
    ("contribution", "Contribution to the Literature and Broader Context"),
    ("background", "Background Details / Context"),
    ("data_sources", "Data Sources"),
    ("concerns_extensions", "Concerns and Extensions"),
    ("additional_notes", "Additional Notes"),
    ("empirical_strategy", "Empirical Strategy and Identification Concerns"),
    ("exhibits", "Tables and Figures"),
]

COMPACT_SECTION_KEYS = [
    "one_minute_summary",
    "contribution",
    "background",
    "data_sources",
    "concerns_extensions",
    "additional_notes",
]


def _text(value: Any) -> str:
    return str(value or "")


def _markdown_text(value: Any) -> str:
    return _text(value)


def _meaningful_source(source: Dict[str, Any]) -> bool:
    return bool(_text(source.get("name")).strip() or _text(source.get("description")).strip())


def _meaningful_exhibit(exhibit: Dict[str, Any]) -> bool:
    return any(
        [
            _text(exhibit.get("identifier_title")).strip(),
            _text(exhibit.get("convey")).strip(),
            _text(exhibit.get("empirical_strategy")).strip(),
            _text(exhibit.get("interpretation")).strip(),
            _text(exhibit.get("qualifications")).strip(),
            exhibit.get("image"),
            exhibit.get("images"),
        ]
    )


def _exhibit_images(exhibit: Dict[str, Any]) -> list[Dict[str, Any]]:
    parts = [part for part in exhibit.get("images", []) if isinstance(part, dict) and part.get("image")]
    if parts:
        return parts
    if exhibit.get("image"):
        return [
            {
                "label": None,
                "page_number": exhibit.get("page_number", ""),
                "image": exhibit["image"],
                "crop": exhibit.get("crop"),
            }
        ]
    return []


def _meaningful_strategy(strategy: Dict[str, Any]) -> bool:
    return any(
        [
            _text(strategy.get("title")).strip(),
            _text(strategy.get("description")).strip(),
            _text(strategy.get("latex_equation")).strip(),
            strategy.get("image"),
        ]
    )


def safe_export_stem(title: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9() ._'-]+", "_", title).strip(" ._")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return (cleaned[:100].rstrip(" ._") or "paper") + "_notes"


def safe_combined_export_stem(title: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9() ._'-]+", "_", title).strip(" ._")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned[:100].rstrip(" ._") or "Combined summaries"


def _unique_path(folder: Path, stem: str, suffix: str) -> Path:
    candidate = folder / f"{stem}{suffix}"
    counter = 2
    while candidate.exists():
        candidate = folder / f"{stem}_{counter}{suffix}"
        counter += 1
    return candidate


def _set_font(style, name: str, size: float) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    for theme_attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(theme_attribute), None)


def _configure_word_document(document: Document, compact: bool = False) -> None:
    sizes = {
        "Normal": 10 if compact else 11,
        "Title": 18,
        "Heading 1": 13,
        "Heading 2": 11.5,
        "Heading 3": 10.5,
        "List Bullet": 10 if compact else 11,
    }
    for style_name, size in sizes.items():
        style = document.styles[style_name]
        _set_font(style, "Arial", size)
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    normal = document.styles["Normal"].paragraph_format
    normal.space_after = Pt(2 if compact else 5)
    normal.line_spacing = 1.0 if compact else 1.08
    for style_name, before, after in [
        ("Title", 0, 6),
        ("Heading 1", 8, 3),
        ("Heading 2", 6, 2),
        ("Heading 3", 3, 1),
    ]:
        paragraph_format = document.styles[style_name].paragraph_format
        paragraph_format.space_before = Pt(before)
        paragraph_format.space_after = Pt(after)
        paragraph_format.keep_with_next = True
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.55 if compact else 0.7)
        section.bottom_margin = Inches(0.55 if compact else 0.7)
        section.left_margin = Inches(0.6 if compact else 0.75)
        section.right_margin = Inches(0.6 if compact else 0.75)


def _set_columns(section, count: int) -> None:
    columns = section._sectPr.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        section._sectPr.append(columns)
    columns.set(qn("w:num"), str(count))
    columns.set(qn("w:space"), "288")


def _finalize_arial(document: Document) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:ascii"), "Arial")
            fonts.set(qn("w:hAnsi"), "Arial")
            fonts.set(qn("w:eastAsia"), "Arial")


def section_has_content(notes: Dict[str, Any], key: str) -> bool:
    sections = notes.get("sections", {})
    if key == "data_sources":
        return any(_meaningful_source(source) for source in notes.get("data_sources", []))
    if key == "empirical_strategy":
        return bool(_text(sections.get(key)).strip()) or any(
            _meaningful_strategy(strategy) for strategy in notes.get("empirical_strategies", [])
        )
    if key == "exhibits":
        return any(_meaningful_exhibit(exhibit) for exhibit in notes.get("exhibits", []))
    return bool(_text(sections.get(key)).strip())


def paper_has_selected_content(notes: Dict[str, Any], selected_sections: Iterable[str]) -> bool:
    return any(section_has_content(notes, key) for key in selected_sections)


def build_markdown(notes: Dict[str, Any]) -> str:
    lines = [f"# {_text(notes.get('title'))}", ""]
    original = _text(notes.get("original_pdf_name"))
    if original:
        lines.extend([f"**Source PDF:** {original}", ""])

    sections = notes.get("sections", {})
    for key, label in SECTION_LABELS:
        if key == "data_sources":
            sources = [source for source in notes.get("data_sources", []) if _meaningful_source(source)]
            if not sources:
                continue
            lines.extend([f"## {label}", ""])
            for index, source in enumerate(sources, start=1):
                name = _text(source.get("name")).strip() or f"Source {index}"
                lines.extend([f"### {name}", ""])
                if _text(source.get("description")).strip():
                    lines.extend([_markdown_text(source.get("description")), ""])
        elif key == "empirical_strategy":
            overview = _text(sections.get(key))
            strategies = [
                strategy
                for strategy in notes.get("empirical_strategies", [])
                if _meaningful_strategy(strategy)
            ]
            if not overview.strip() and not strategies:
                continue
            lines.extend([f"## {label}", ""])
            if overview.strip():
                lines.extend([overview, ""])
            for index, strategy in enumerate(strategies, start=1):
                title = _text(strategy.get("title")).strip() or f"Empirical Strategy {index}"
                lines.extend([f"### {title}", ""])
                if strategy.get("image"):
                    lines.extend([f"![{title}](../{strategy['image']})", ""])
                lines.extend([f"**Page:** {strategy.get('page_number', '')}", ""])
                if _text(strategy.get("description")).strip():
                    lines.extend(["**Description**", "", _text(strategy.get("description")), ""])
                if _text(strategy.get("latex_equation")).strip():
                    lines.extend(
                        ["**Equation**", "", "$$", _text(strategy.get("latex_equation")), "$$", ""]
                    )
        elif key == "exhibits":
            exhibits = [exhibit for exhibit in notes.get("exhibits", []) if _meaningful_exhibit(exhibit)]
            if not exhibits:
                continue
            lines.extend([f"## {label}", ""])
            for index, exhibit in enumerate(exhibits, start=1):
                title = _text(exhibit.get("identifier_title")).strip() or f"Exhibit {index}"
                lines.extend([f"### {title}", ""])
                image_parts = _exhibit_images(exhibit)
                for part in image_parts:
                    label = _text(part.get("label")).replace("_", " ").strip()
                    alt = f"{title} — {label}" if label else title
                    lines.extend([f"![{alt}](../{part['image']})", ""])
                pages = [str(part.get("page_number", "")) for part in image_parts]
                page_text = ", ".join(dict.fromkeys(page for page in pages if page))
                if not page_text:
                    page_text = str(exhibit.get("page_number", ""))
                lines.extend([f"**Page{'s' if len(set(pages)) > 1 else ''}:** {page_text}", ""])
                fields = [
                    ("What it is trying to convey", "convey"),
                    ("Empirical strategy used", "empirical_strategy"),
                    ("Interpretation / notes", "interpretation"),
                    ("Qualifications / concerns", "qualifications"),
                ]
                for heading, field in fields:
                    if _text(exhibit.get(field)).strip():
                        lines.extend([f"**{heading}**", "", _markdown_text(exhibit.get(field)), ""])
        else:
            authored = _text(sections.get(key))
            if not authored.strip():
                continue
            lines.extend([f"## {label}", "", authored, ""])
    return "\n".join(lines).rstrip() + "\n"


def export_markdown(data_root: Path, notes: Dict[str, Any]) -> Path:
    folder = record_dir(data_root, str(notes["record_id"]))
    output = folder / "formatted_notes" / f"{safe_export_stem(_text(notes.get('title')))}.md"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(build_markdown(notes), encoding="utf-8")
    return output


def _add_authored_text(document: Document, value: Any) -> None:
    text = _text(value)
    if not text.strip():
        return
    for line in text.splitlines() or [""]:
        stripped = line.lstrip()
        if stripped.startswith(("- ", "* ")):
            document.add_paragraph(stripped[2:], style="List Bullet")
        else:
            document.add_paragraph(line)


def export_docx(data_root: Path, notes: Dict[str, Any]) -> Path:
    folder = record_dir(data_root, str(notes["record_id"]))
    output = folder / "formatted_notes" / f"{safe_export_stem(_text(notes.get('title')))}.docx"
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_word_document(document, compact=True)
    document.core_properties.title = _text(notes.get("title"))
    document.add_heading(_text(notes.get("title")), level=0)
    if notes.get("original_pdf_name"):
        document.add_paragraph(f"Source PDF: {notes['original_pdf_name']}")

    has_compact = any(section_has_content(notes, key) for key in COMPACT_SECTION_KEYS)
    has_empirical = section_has_content(notes, "empirical_strategy")
    has_figures = section_has_content(notes, "exhibits")

    working_section = document.add_section(WD_SECTION.CONTINUOUS)
    _set_columns(working_section, 2 if has_compact or has_empirical else 1)
    if has_compact:
        _add_individual_review(document, notes)
    if has_empirical:
        if has_compact:
            working_section = document.add_section(WD_SECTION.NEW_PAGE)
            _set_columns(working_section, 2)
        document.add_heading("Empirical Strategy and Identification Concerns", level=1)
        _add_empirical_strategy_for_paper(document, data_root, notes, heading_level=2)
    if has_figures:
        if has_compact or has_empirical:
            working_section = document.add_section(WD_SECTION.NEW_PAGE)
        _set_columns(working_section, 1)
        document.add_heading("Tables and Figures", level=1)
        _add_figures_for_paper(document, data_root, notes, heading_level=2)

    _finalize_arial(document)
    document.save(output)
    return output


def export_combined_summaries(
    data_root: Path,
    selected_notes: Iterable[Dict[str, Any]],
    export_title: str,
) -> Path:
    return export_combined_papers(
        data_root,
        selected_notes,
        export_title,
        selected_sections=["one_minute_summary"],
    )


def _add_categories(document: Document, notes: Dict[str, Any]) -> None:
    categories = [str(item).strip() for item in notes.get("categories", []) if str(item).strip()]
    if categories:
        paragraph = document.add_paragraph()
        paragraph.add_run("Categories: ").bold = True
        paragraph.add_run(", ".join(categories))


def _add_data_sources(document: Document, notes: Dict[str, Any]) -> None:
    for index, source in enumerate(notes.get("data_sources", []), start=1):
        if not _meaningful_source(source):
            continue
        paragraph = document.add_paragraph()
        name = _text(source.get("name")).strip() or f"Source {index}"
        paragraph.add_run(name).bold = True
        description = _text(source.get("description")).strip()
        if description:
            paragraph.add_run(": " + description)


def _add_individual_review(document: Document, notes: Dict[str, Any]) -> None:
    labels = dict(GROUP_EXPORT_SECTIONS)
    for key in COMPACT_SECTION_KEYS:
        if not section_has_content(notes, key):
            continue
        document.add_heading(labels[key], level=1)
        if key == "data_sources":
            _add_data_sources(document, notes)
        else:
            _add_authored_text(document, notes.get("sections", {}).get(key))


def _add_compact_review(
    document: Document,
    papers: list[Dict[str, Any]],
    selected_sections: set[str],
) -> None:
    labels = dict(GROUP_EXPORT_SECTIONS)
    document.add_heading("Review Notes", level=1)
    for notes in papers:
        available = [
            key
            for key in COMPACT_SECTION_KEYS
            if key in selected_sections and section_has_content(notes, key)
        ]
        if not available:
            continue
        document.add_heading(_text(notes.get("title")), level=2)
        _add_categories(document, notes)
        for key in available:
            document.add_heading(labels[key], level=3)
            if key == "data_sources":
                _add_data_sources(document, notes)
            else:
                _add_authored_text(document, notes.get("sections", {}).get(key))


def _add_empirical_strategy_for_paper(
    document: Document,
    data_root: Path,
    notes: Dict[str, Any],
    heading_level: int,
) -> None:
    overview = _text(notes.get("sections", {}).get("empirical_strategy"))
    if overview.strip():
        _add_authored_text(document, overview)
    folder = record_dir(data_root, str(notes["record_id"]))
    for index, strategy in enumerate(notes.get("empirical_strategies", []), start=1):
        if not _meaningful_strategy(strategy):
            continue
        title = _text(strategy.get("title")).strip() or f"Empirical Strategy {index}"
        document.add_heading(title, level=heading_level)
        if strategy.get("image"):
            image_path = folder / str(strategy["image"])
            if image_path.is_file():
                document.add_picture(str(image_path), width=Inches(2.75))
        document.add_paragraph(f"Page: {strategy.get('page_number', '')}")
        if _text(strategy.get("description")).strip():
            _add_authored_text(document, strategy.get("description"))
        if _text(strategy.get("latex_equation")).strip():
            paragraph = document.add_paragraph()
            paragraph.add_run("Equation: ").bold = True
            paragraph.add_run(_text(strategy.get("latex_equation")))


def _add_empirical_strategies(document: Document, data_root: Path, papers: list[Dict[str, Any]]) -> None:
    document.add_heading("Empirical Strategy and Identification Concerns", level=1)
    for notes in papers:
        if not section_has_content(notes, "empirical_strategy"):
            continue
        document.add_heading(_text(notes.get("title")), level=2)
        _add_empirical_strategy_for_paper(document, data_root, notes, heading_level=3)


def _add_figures_for_paper(
    document: Document,
    data_root: Path,
    notes: Dict[str, Any],
    heading_level: int,
) -> None:
    exhibits = [item for item in notes.get("exhibits", []) if _meaningful_exhibit(item)]
    folder = record_dir(data_root, str(notes["record_id"]))
    for index, exhibit in enumerate(exhibits, start=1):
        title = _text(exhibit.get("identifier_title")).strip() or f"Exhibit {index}"
        document.add_heading(title, level=heading_level)
        image_parts = _exhibit_images(exhibit)
        for part in image_parts:
            label = _text(part.get("label")).replace("_", " ").strip()
            if label:
                document.add_paragraph(label.replace("panel ", "Panel ", 1))
            image_path = folder / str(part["image"])
            if image_path.is_file():
                document.add_picture(str(image_path), width=Inches(6.2))
        pages = [str(part.get("page_number", "")) for part in image_parts]
        page_text = ", ".join(dict.fromkeys(page for page in pages if page))
        if not page_text:
            page_text = str(exhibit.get("page_number", ""))
        document.add_paragraph(f"Page{'s' if len(set(pages)) > 1 else ''}: {page_text}")
        for heading, field in [
            ("What it is trying to convey", "convey"),
            ("Empirical strategy used", "empirical_strategy"),
            ("Interpretation / notes", "interpretation"),
            ("Qualifications / concerns", "qualifications"),
        ]:
            value = _text(exhibit.get(field))
            if value.strip():
                paragraph = document.add_paragraph()
                paragraph.add_run(heading + ": ").bold = True
                paragraph.add_run(value)


def _add_figures(document: Document, data_root: Path, papers: list[Dict[str, Any]]) -> None:
    document.add_heading("Tables and Figures", level=1)
    for notes in papers:
        if not section_has_content(notes, "exhibits"):
            continue
        document.add_heading(_text(notes.get("title")), level=2)
        _add_figures_for_paper(document, data_root, notes, heading_level=3)


def export_combined_papers(
    data_root: Path,
    selected_notes: Iterable[Dict[str, Any]],
    export_title: str,
    selected_sections: Iterable[str],
) -> Path:
    selected = set(selected_sections)
    allowed = {key for key, _ in GROUP_EXPORT_SECTIONS}
    selected &= allowed
    if not selected:
        raise ValueError("Choose at least one section to export.")
    papers = order_records(
        data_root,
        [notes for notes in selected_notes if paper_has_selected_content(notes, selected)],
    )
    if not papers:
        raise ValueError("None of the selected papers has content in the chosen sections.")

    output_folder = data_root.resolve().parent / "combined_exports"
    output_folder.mkdir(parents=True, exist_ok=True)
    output = _unique_path(output_folder, safe_combined_export_stem(export_title), ".docx")

    document = Document()
    _configure_word_document(document, compact=True)
    document.core_properties.title = export_title.strip() or "Combined Paper Notes"
    document.add_heading(export_title.strip() or "Combined Paper Notes", level=0)
    document.add_paragraph(f"Created {datetime.now().strftime('%B %d, %Y')}")

    has_compact = any(
        key in selected and any(section_has_content(notes, key) for notes in papers)
        for key in COMPACT_SECTION_KEYS
    )
    has_empirical = "empirical_strategy" in selected and any(
        section_has_content(notes, "empirical_strategy") for notes in papers
    )
    has_figures = "exhibits" in selected and any(
        section_has_content(notes, "exhibits") for notes in papers
    )

    first_column_count = 2 if has_compact or has_empirical else 1
    working_section = document.add_section(WD_SECTION.CONTINUOUS)
    _set_columns(working_section, first_column_count)
    if has_compact:
        _add_compact_review(document, papers, selected)
    if has_empirical:
        if has_compact:
            working_section = document.add_section(WD_SECTION.NEW_PAGE)
            _set_columns(working_section, 2)
        _add_empirical_strategies(document, data_root, papers)
    if has_figures:
        if has_compact or has_empirical:
            working_section = document.add_section(WD_SECTION.NEW_PAGE)
        _set_columns(working_section, 1)
        _add_figures(document, data_root, papers)

    _finalize_arial(document)
    document.save(output)
    return output
